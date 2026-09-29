import json
import os
import openpyxl
from datetime import datetime

def load_mapping(config_path):
    """Loads the De-Para mapping from a JSON file."""
    if os.path.exists(config_path):
        try:
            with open(config_path, 'r', encoding='utf-8') as f:
                return json.load(f)
        except Exception:
            return {}
    return {}

def save_mapping(mapping, config_path):
    """Saves the De-Para mapping to a JSON file."""
    try:
        os.makedirs(os.path.dirname(config_path), exist_ok=True)
        with open(config_path, 'w', encoding='utf-8') as f:
            json.dump(mapping, f, ensure_ascii=False, indent=4)
        return True
    except Exception as e:
        print(f"Error saving mapping: {str(e)}")
        return False

def generate_entries(parsed_data, mapping, batch_number="1", entry_date=None, doc_number="FOLHA"):
    """
    Generates accounting entry rows based on parsed events and user mapping.
    
    mapping structure:
    {
        "event_code": {
            "debit_account": str,
            "credit_account": str,
            "cost_center": str
        }
    }
    """
    events = list(parsed_data.get('events') or [])
    meta = parsed_data.get('metadata') or {}
    empresa_code = meta.get('empresa_codigo', '0001')
    
    # Process social charges as virtual events if present
    sc = parsed_data.get('social_charges', {})
    if sc:
        patronal_total = sc.get('gps_empresa_func', 0.0) + sc.get('gps_empresa_socios', 0.0) + sc.get('gps_empresa_auton', 0.0)
        if abs(patronal_total) >= 0.01:
            events.append({
                'code': 'GPS_PATRONAL',
                'description': 'INSS Patronal (Empresa)',
                'total': patronal_total,
                'section': 'social_charges',
                'type': 'encargo'
            })
        if abs(sc.get('gps_rat', 0.0)) >= 0.01:
            events.append({
                'code': 'GPS_RAT',
                'description': 'INSS RAT/FAP',
                'total': sc['gps_rat'],
                'section': 'social_charges',
                'type': 'encargo'
            })
        if abs(sc.get('gps_terceiros', 0.0)) >= 0.01:
            events.append({
                'code': 'GPS_TERCEIROS',
                'description': 'INSS Terceiros',
                'total': sc['gps_terceiros'],
                'section': 'social_charges',
                'type': 'encargo'
            })
        if abs(sc.get('fgts_total', 0.0)) >= 0.01:
            events.append({
                'code': 'FGTS',
                'description': 'FGTS (Total Apurado)',
                'total': sc['fgts_total'],
                'section': 'social_charges',
                'type': 'encargo'
            })
    
    # Definir centro de custo padrão fixo por empresa (0001 -> 1SVP, 0002 -> 2MCJ, 0003 -> 3FAB)
    STANDARD_COMPANY_CC = {
        "0001": "1SVP",
        "0002": "2MCJ",
        "0003": "3FAB"
    }
    company_fixed_cc = STANDARD_COMPANY_CC.get(str(empresa_code).strip(), "1SVP")
        
    # Format month/year for history (e.g. 05/2026 -> MAI/26)
    mes_ano = ""
    if meta.get('periodo_referencia'):
        try:
            parts = meta['periodo_referencia'].split('/')
            if len(parts) == 2:
                mes, ano = parts[0], parts[1]
                meses_pt = {
                    "01": "JAN", "02": "FEV", "03": "MAR", "04": "ABR",
                    "05": "MAI", "06": "JUN", "07": "JUL", "08": "AGO",
                    "09": "SET", "10": "OUT", "11": "NOV", "12": "DEZ"
                }
                mes_abrev = meses_pt.get(mes, mes)
                ano_abrev = ano[-2:]
                mes_ano = f"{mes_abrev}/{ano_abrev}"
        except Exception:
            pass
    if not mes_ano:
        mes_ano = datetime.now().strftime("%b/%y").upper()
        
    if not entry_date:
        # Default to the end date of the period or today
        if meta.get('periodo_fim'):
            # Assume DD/MM/YYYY
            try:
                date_obj = datetime.strptime(meta['periodo_fim'], "%d/%m/%Y")
                entry_date = date_obj.strftime("%d%m%Y")
            except ValueError:
                entry_date = datetime.now().strftime("%d%m%Y")
        else:
            entry_date = datetime.now().strftime("%d%m%Y")
    else:
        # Standardize date format to DDMMAAAA
        entry_date = entry_date.replace("/", "").replace("-", "")
        
    rows = []
    sequence = 1
    unmapped_events = []
    
    total_debit = 0.0
    total_credit = 0.0
    
    # Process each event
    for ev in events:
        code = ev['code']
        desc = ev['description']
        amount = ev['total']
        
        # Skip events with zero value
        if abs(amount) < 0.01:
            continue
            
        ev_map = mapping.get(code, {})
        debit_acc = ev_map.get('debit_account', '').strip()
        credit_acc = ev_map.get('credit_account', '').strip()
        cc = ev_map.get('cost_center', '').strip()
        
        # Track unmapped events (if both accounts are missing, or at least one is needed)
        # Note: some events might intentionally map only to Debit or only to Credit in a complex entry,
        # but in a simple event-by-event double entry, we need both.
        if not debit_acc or not credit_acc:
            unmapped_events.append({
                'code': code,
                'description': desc,
                'amount': amount,
                'debit_missing': not debit_acc,
                'credit_missing': not credit_acc
            })
            
        # Leg 1: Debit Row
        if debit_acc:
            clean_debit = debit_acc.replace(".", "").strip()
            use_cc_debit = (company_fixed_cc if company_fixed_cc else cc) if (clean_debit.startswith('3') or clean_debit.startswith('4') or clean_debit.startswith('5')) else None
            
            # Format history: DESCRIÇÃO CÓD DA FOLHA + FOLHA + CENTRO DE CUSTO + MES E ANO in uppercase
            cc_suffix = f" {use_cc_debit}" if use_cc_debit else ""
            hist_debit = f"{desc}-FOLHA{cc_suffix} - {mes_ano}".upper()
            hist_debit = hist_debit[:200]
            
            rows.append({
                'empresa': empresa_code,
                'lote': batch_number,
                'data': entry_date,
                'documento': doc_number,
                'conta': debit_acc,
                'cc': use_cc_debit if use_cc_debit else None,
                'tipo': 'D',
                'historico': hist_debit,
                'valor': amount,
                'sequencia': sequence
            })
            total_debit += amount
            sequence += 1
            
        # Leg 2: Credit Row
        if credit_acc:
            clean_credit = credit_acc.replace(".", "").strip()
            use_cc_credit = (company_fixed_cc if company_fixed_cc else cc) if (clean_credit.startswith('3') or clean_credit.startswith('4') or clean_credit.startswith('5')) else None
            
            # Format history: DESCRIÇÃO CÓD DA FOLHA + FOLHA + CENTRO DE CUSTO + MES E ANO in uppercase
            cc_suffix = f" {use_cc_credit}" if use_cc_credit else ""
            hist_credit = f"{desc}-FOLHA{cc_suffix} - {mes_ano}".upper()
            hist_credit = hist_credit[:200]
            
            rows.append({
                'empresa': empresa_code,
                'lote': batch_number,
                'data': entry_date,
                'documento': doc_number,
                'conta': credit_acc,
                'cc': use_cc_credit if use_cc_credit else None,
                'tipo': 'C',
                'historico': hist_credit,
                'valor': amount,
                'sequencia': sequence
            })
            total_credit += amount
            sequence += 1
            
    is_balanced = abs(total_debit - total_credit) < 0.01
    
    return {
        'rows': rows,
        'total_debit': total_debit,
        'total_credit': total_credit,
        'difference': abs(total_debit - total_credit),
        'is_balanced': is_balanced,
        'unmapped_events': unmapped_events
    }

def write_to_excel_template(template_path, output_path, entries_data):
    """
    Loads the MXM template Excel file, populates the 'Dados' sheet with generated entries,
    and saves the workbook to output_path.
    """
    if not os.path.exists(template_path):
        raise FileNotFoundError(f"Template Excel file not found: {template_path}")
        
    wb = openpyxl.load_workbook(template_path)
    
    # Ensure sheet 'Dados' exists
    if 'Dados' not in wb.sheetnames:
        raise ValueError("Sheet 'Dados' not found in the Excel template.")
        
    sheet = wb['Dados']
    
    # Clear existing data rows under header
    # Header is on row 1. Let's delete all rows from row 2 onwards.
    if sheet.max_row > 1:
        sheet.delete_rows(2, sheet.max_row)
        
    # Columns mapping in the sheet 'Dados' based on the template:
    # A (1): Codigo da empresa
    # B (2): Lote
    # C (3): Data do lancamento
    # D (4): Documento
    # E (5): Conta contabil
    # F (6): Centro de Custo
    # G (7): Indicador de Debito ou Credito
    # H (8): Historico do Lancamento
    # I (9): Blank/Unnamed: 8
    # J (10): Valor do Lancamento na 2a Moeda (blank)
    # K (11): Codigo da Segunda Moeda (blank)
    # L (12): Valor do Lancamento (numeric)
    # M (13): Sequencia do Lancamento (numeric)
    # N (14): Codigo do projeto (blank)
    # O (15): Codigo do Fornecedor (blank)
    # P (16): Codigo do Cliente (blank)
    # Q (17): Numero do Titulo (blank)
    # R (18): Converter Moeda (blank)
    # S (19): Historico Padrao (blank)
    # T (20): Hist. Padrao - Complemento 1 (blank)
    # U (21): Hist. Padrao - Complemento 2 (blank)
    # V (22): Hist. Padrao - Complemento 3 (blank)
    # W (23): Excluir Lancamentos (blank)

    for entry in entries_data:
        row_idx = sheet.max_row + 1
        
        sheet.cell(row=row_idx, column=1, value=str(entry['empresa'])[:4])
        sheet.cell(row=row_idx, column=2, value=str(entry['lote'])[:6])
        sheet.cell(row=row_idx, column=3, value=str(entry['data'])[:8])
        sheet.cell(row=row_idx, column=4, value=str(entry['documento'])[:6] if entry['documento'] else "")
        sheet.cell(row=row_idx, column=5, value=str(entry['conta'])[:15])
        sheet.cell(row=row_idx, column=6, value=str(entry['cc'])[:15] if entry['cc'] else "")
        sheet.cell(row=row_idx, column=7, value=str(entry['tipo'])[:1])
        sheet.cell(row=row_idx, column=8, value=str(entry['historico'])[:200])
        
        # Column I (9) is blank
        sheet.cell(row=row_idx, column=9, value="")
        # Column J (10) 2nd currency is blank
        sheet.cell(row=row_idx, column=10, value="")
        # Column K (11) 2nd currency code is blank
        sheet.cell(row=row_idx, column=11, value="")
        
        # Column L (12) is Valor do Lancamento
        sheet.cell(row=row_idx, column=12, value=float(entry['valor']))
        
        # Column M (13) is Sequencia
        sheet.cell(row=row_idx, column=13, value=int(entry['sequencia']))
        
        # Fill rest of columns as blank or empty strings
        for c in range(14, 24):
            sheet.cell(row=row_idx, column=c, value="")
            
    wb.save(output_path)
    return True
